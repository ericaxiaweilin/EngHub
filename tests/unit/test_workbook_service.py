from pathlib import Path

from openpyxl import Workbook, load_workbook

from api.services.workbook_service import (
    apply_workbook_operations,
    build_pmc_workbook_template,
    build_pivot_summary,
    snapshot_to_table,
    workbook_snapshot_to_xlsx,
    xlsx_to_workbook_snapshot,
)


def test_xlsx_roundtrip_preserves_sheets_formulas_and_freeze(tmp_path: Path):
    source = tmp_path / "pmc.xlsx"
    output = tmp_path / "roundtrip.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "MRP"
    sheet.append(["物料", "毛需求", "库存", "净需求"])
    sheet.append(["MAT-001", 100, 40, "=B2-C2"])
    sheet.freeze_panes = "A2"
    bom = book.create_sheet("BOM")
    bom["A1"] = "产品"
    bom["B1"] = "物料"
    book.save(source)

    snapshot = xlsx_to_workbook_snapshot(source)
    mrp = next(value for value in snapshot["sheets"].values() if value["name"] == "MRP")
    assert len(snapshot["sheetOrder"]) == 2
    assert mrp["cellData"]["1"]["3"]["f"] == "=B2-C2"
    assert mrp["freeze"] == {"xSplit": 0, "ySplit": 1}

    snapshot, changed = apply_workbook_operations(snapshot, [
        {"type": "set_formula", "sheet": "MRP", "cell": "E2", "formula": "=SUM(B2:C2)"},
        {"type": "append_rows", "sheet": "BOM", "rows": [["P-001", "MAT-001"]]},
    ])
    assert len(changed) == 2
    assert snapshot_to_table(snapshot, "MRP")["formula_cells"] == 2

    workbook_snapshot_to_xlsx(snapshot, output)
    exported = load_workbook(output, data_only=False)
    assert exported.sheetnames == ["MRP", "BOM"]
    assert exported["MRP"]["D2"].value == "=B2-C2"
    assert exported["MRP"]["E2"].value == "=SUM(B2:C2)"
    assert exported["BOM"]["A2"].value == "P-001"


def test_pmc_template_contains_formula_rows(tmp_path: Path):
    snapshot = build_pmc_workbook_template()
    mrp = snapshot["sheets"]["pmc-mrp"]
    assert mrp["cellData"]["0"]["0"]["v"] == "物料编码"
    assert mrp["cellData"]["1"]["5"]["f"] == '=IF(A2="","",MAX(C2+E2-D2,0))'
    output = tmp_path / "pmc-template.xlsx"
    workbook_snapshot_to_xlsx(snapshot, output)
    reopened = load_workbook(output, data_only=False)
    assert reopened["MRP"]["F2"].value == '=IF(A2="","",MAX(C2+E2-D2,0))'
    assert reopened.calculation.fullCalcOnLoad is True


def test_flash_fill_and_pivot_summary():
    snapshot = {
        "id": "test",
        "sheetOrder": ["data"],
        "sheets": {
            "data": {
                "id": "data",
                "name": "明细",
                "rowCount": 20,
                "columnCount": 10,
                "cellData": {
                    "0": {"0": {"v": "客户"}, "1": {"v": "金额"}, "2": {"v": "编码"}},
                    "1": {"0": {"v": "A"}, "1": {"v": 10}, "2": {"v": "A1001-2024-01"}, "3": {"v": "2024"}},
                    "2": {"0": {"v": "A"}, "1": {"v": 20}, "2": {"v": "A1002-2024-02"}, "3": {"v": "2024"}},
                    "3": {"0": {"v": "B"}, "1": {"v": 5}, "2": {"v": "B1003-2025-01"}},
                },
            },
        },
    }
    snapshot, changed = apply_workbook_operations(snapshot, [{
        "type": "flash_fill", "sheet": "明细", "source_column": "C", "target_column": "D", "start_row": 2, "end_row": 4,
    }])
    assert changed[0]["filled"] == 1
    assert snapshot["sheets"]["data"]["cellData"]["3"]["3"]["v"] == "2025"

    snapshot, summary = build_pivot_summary(snapshot, "明细", "客户", "金额", "sum")
    assert summary["groups"] == 2
    pivot = next(sheet for sheet in snapshot["sheets"].values() if sheet["name"] == "透视汇总")
    assert pivot["cellData"]["1"]["0"]["v"] == "A"
    assert pivot["cellData"]["1"]["1"]["v"] == 30
